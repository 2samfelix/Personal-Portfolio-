#!/usr/bin/env python3
"""Export workbook previews to the JSON read by WorkbookPreview.

Offline tool, run by hand whenever a workbook changes:

    npm run export:model-preview              # every project
    python3 scripts/export-model-preview.py arm

Requires Python 3 with openpyxl, and LibreOffice (`soffice`) on PATH.

- openpyxl supplies structure, formulas, cached (raw) values and styles.
- LibreOffice's CSV export "as shown" supplies the displayed text, so number
  formats render exactly as the spreadsheet shows them (no formatter of our own).

A project exports either whole sheets (one tab per sheet) or curated views (one
tab per view: a sheet plus one or more ranges, with a caption, and optionally
`freeze_cols` to keep its label columns in view). Views keep the real row and
column numbers, so the formula bar shows true cell references.

The output is a read-only snapshot: nothing is recalculated in the browser. The
workbook's SHA-256 is recorded so `npm run check:model-preview` can detect a
workbook that changed without the JSON being regenerated.

A workbook saved without cached formula results (project option `recalculate`)
is recalculated in a temporary copy by LibreOffice: formulas, styles, widths and
heights still come from the original, values and displayed text from the copy.
Every exported formula must read the same in both, and every recalculated value
must exist and must not be an error. The original is never modified.

Colour-scale conditional formats (2- and 3-colour; min, max, number, percent and
percentile thresholds) are evaluated here, as Excel does, and stored as each
cell's fill; the preview itself never evaluates rules.

Fails loudly (non-zero exit) if a requested sheet is missing, a formula has no
cached value (other than known empty-string text formulas), a workbook copy does
not match its recorded source hash, or the sheets use something this preview
cannot show faithfully (tinted theme or indexed colours, conditional formats
other than colour scales, a merged range cut by a view's range).
"""
from __future__ import annotations

import csv
import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries

ROOT = Path(__file__).resolve().parent.parent

PROJECTS: dict[str, dict] = {
    "caterpillar": {
        "workbook": "caterpillar-demand-quality/model/CAT_8Q_Operating_Dataset.xlsx",
        "output": "public/caterpillar/model-preview.json",
        # Story sheets, in tab order.
        "sheets": ["Analysis", "Construction Industries", "Dealer Inventory", "Retail Statistics"],
        # Text formulas of the form IF(x="","",x) whose result is an empty string. The
        # file stores no cached value for them, which is correct, so they are shown as
        # empty cells (the formula is kept so the formula bar can still show it).
        "known_empty_text_formulas": {
            ("Analysis", a) for a in ["B15", "C15", "D15", "E15", "F15", "G15", "B16", "E16", "F16", "I16"]
        },
    },
    "arm": {
        # A copy of the workbook from the Arm repository; the export refuses to run
        # if the copy no longer matches the source hash recorded here.
        "workbook": "arm-ipo-valuation/model/Sam_Felix_Arm_IPO_Valuation_Strategy_Analysis (7).xlsx",
        "output": "public/arm/model-preview.json",
        "source": {
            "repo": "2samfelix/arm-ipo-valuation-strategy-analysis",
            "commit": "24f1d107753f59a08eaccc8f581057eea4beab5a",
            "path": "Sam_Felix_Arm_IPO_Valuation_Strategy_Analysis (7).xlsx",
            "sha256": "ce1574355f5bc59bdffd4a6ead199e24a455d2e327649c9c5a22485c7413d47e",
        },
        "views": [
            {
                "name": "IPO vs. peers",
                "sheet": "IPO Valuation",
                "ranges": ["A15:F19"],
                "caption": "Arm priced above every central peer-implied value.",
                "freeze_cols": 1,
            },
            {
                "name": "From WACC to share price",
                "sheet": "DCF Valuation",
                "ranges": ["A4:B11", "A42:B50"],
                "caption": "The DCF-implied price is about $51.54, against a $272.21 market price on Aug 10, 2026.",
                "freeze_cols": 1,
            },
            {
                "name": "Sensitivity",
                "sheet": "DCF Valuation",
                "ranges": ["A52:F58"],
                "caption": "Every cell is below the current price. Highest value in the grid: $94.39.",
                "freeze_cols": 1,
            },
        ],
    },
    "microsoft-activision": {
        # A copy of the workbook from the Microsoft x Activision repository; the export
        # refuses to run if the copy no longer matches the source hash recorded here.
        "workbook": "microsoft-activision-ma/model/Sam_Felix_MSFT_ATVI_MA_Model.xlsx",
        "output": "public/microsoft-activision/model-preview.json",
        "source": {
            "repo": "2samfelix/microsoft-activision-ma-analysis",
            "commit": "a5ed679f114caef98fc5fb60c1171e109c0a90b4",
            "path": "Sam_Felix_MSFT_ATVI_MA_Model.xlsx",
            "sha256": "3a5106c1ea0e699f05aa0e186a9dcc1b9e42affb9d542414e8a1d71bcde09340",
        },
        "views": [
            {
                "name": "Why the deal dilutes",
                "sheet": "Accretion Dilution",
                "ranges": ["A12:D22", "A25:D28"],
                "caption": "With zero synergies, pro forma EPS is $9.29 against $9.68 standalone: 4.07% dilution. My base case sits $0.26 below Microsoft's reported pro forma EPS.",
                "freeze_cols": 1,
            },
            {
                "name": "How much synergy is enough?",
                "sheet": "Synergy Breakeven",
                "ranges": ["A5:D11"],
                "caption": "Breakeven needs about $3.64B of annual pre-tax synergies, roughly 48% of Activision's FY2022 revenue.",
                "freeze_cols": 1,
            },
            {
                "name": "Sensitivity",
                "sheet": "Synergy Breakeven",
                "ranges": ["A19:F26"],
                "caption": "Red marks the most dilutive cell (−5.5%), yellow the grid's median (−0.7%), green the most accretive (+4.0%). Accretive: $3B at a 2% yield, $4B up to 4%, $5B+ at any yield.",
                "freeze_cols": 1,
            },
        ],
    },
    "nike": {
        # A copy of the workbook from the Nike repository; the export refuses to run
        # if the copy no longer matches the source hash recorded here.
        "workbook": "nike-three-statement-model/model/Sam_Felix_Nike_Three_Statement_Model (4).xlsx",
        "output": "public/nike/model-preview.json",
        "source": {
            "repo": "2samfelix/nike-three-statement-model",
            "commit": "356bcae65f6be02507865391d0b54b4bf824685d",
            "path": "Sam_Felix_Nike_Three_Statement_Model (4).xlsx",
            "sha256": "1568052d25ed0876610b651c8e24e7afe1d9fa7a090e9f1d6280183245795d04",
        },
        # Saved without cached formula results: values come from a LibreOffice recalculation.
        "recalculate": True,
        "views": [
            {
                "name": "Three scenarios, one model",
                "sheet": "Scenario Dashboard",
                "ranges": ["A4:D11", "A14:D19"],
                "caption": "FY2029E diluted EPS is $3.40 in the bull case, $2.39 in the base case and $0.66 in the bear case.",
                "freeze_cols": 1,
            },
            {
                "name": "Does the model beat the street?",
                "sheet": "Scenario Dashboard",
                "ranges": ["A23:D27"],
                "caption": "The base case is 3.9% above consensus EPS in FY2027E and 9.3% below in FY2029E.",
                "freeze_cols": 1,
            },
            {
                "name": "Does North America offset Greater China?",
                "sheet": "Scenario Dashboard",
                "ranges": ["A33:B42"],
                "caption": "Greater China loses about $1.34B of revenue between FY2026A and FY2028E. North America needs about 3.2% annual growth to offset it; the base case has 4.25%, a cushion of about $439M.",
                "freeze_cols": 1,
            },
        ],
    },
}

# Generous caps for whole-sheet tabs; every included sheet is well inside them. A
# sheet that grows past them is truncated and flagged in the JSON, not silently cut.
MAX_ROWS = 120
MAX_COLS = 16

# Excel's defaults when a sheet does not set its own.
DEFAULT_COL_WIDTH = 8.43
DEFAULT_ROW_HEIGHT = 12.75  # points, for 10pt Arial
# The preview's default font; other fonts are named in the style.
DEFAULT_FONT_NAME = "Arial"
# Height of the spacer row between two ranges of one view, in points.
GAP_ROW_PT = 9

FORMAT_VERSION = 1

# Excel's theme colour index -> the colour's name in xl/theme/theme1.xml.
THEME_SLOTS = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
               "accent4", "accent5", "accent6", "hlink", "folHlink"]
THEME: dict[int, str] = {}  # filled from the workbook by load_theme()

ENTITY = re.compile(r"&(?:[A-Za-z]+|#\d+|#x[0-9A-Fa-f]+);")
DECODED: list[str] = []  # cells whose literal HTML entities were decoded for display
# For a recalculated workbook: each sheet's formulas as the recalculated copy holds them.
RECALC_FORMULAS: dict = {}


def fail(message: str) -> None:
    print(f"export-model-preview: ERROR: {message}", file=sys.stderr)
    sys.exit(1)


def note(message: str) -> None:
    print(f"export-model-preview: note: {message}")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_theme(workbook: Path) -> None:
    """The workbook's theme palette, so plain theme colours resolve to RGB."""
    THEME.clear()
    xml = zipfile.ZipFile(workbook).read("xl/theme/theme1.xml").decode("utf-8")
    for i, slot in enumerate(THEME_SLOTS):
        m = re.search(rf"<a:{slot}>.*?(?:srgbClr val|sysClr [^>]*lastClr)=\"([0-9A-Fa-f]{{6}})\"", xml, re.S)
        if m:
            THEME[i] = f"#{m.group(1).upper()}"


def soffice_path() -> str:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        fail("LibreOffice (soffice) is required and was not found on PATH")
    return soffice


def libreoffice_version() -> str:
    out = subprocess.run([soffice_path(), "--version"], check=True, capture_output=True, text=True, timeout=60)
    return out.stdout.strip()


def recalculated_copy(workbook: Path, tmp: str) -> Path:
    """A copy of the workbook with every formula recalculated by LibreOffice, in tmp.
    The original is only read."""
    src = Path(tmp, "in", workbook.name)
    src.parent.mkdir()
    shutil.copyfile(workbook, src)
    profile = Path(tmp, "profile").as_uri()
    subprocess.run(
        [soffice_path(), "--headless", "--norestore", f"-env:UserInstallation={profile}",
         "--convert-to", "xlsx", "--outdir", str(Path(tmp, "out")), str(src)],
        check=True, capture_output=True, timeout=300,
    )
    out = Path(tmp, "out", workbook.stem + ".xlsx")
    if not out.exists():
        fail(f"LibreOffice did not produce a recalculated copy of {workbook.name}")
    return out


def shown_text(workbook: Path) -> dict[str, list[list[str]]]:
    """Every sheet's cells as LibreOffice displays them, via its CSV export."""
    soffice = soffice_path()
    with tempfile.TemporaryDirectory() as tmp:
        profile = Path(tmp, "profile").as_uri()
        # Filter tokens: comma, double quote, UTF-8, from line 1, -, -, quoted text off,
        # detect numbers on, save cell contents AS SHOWN on, formulas off, -, all sheets.
        csv_filter = "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true,false,false,-1"
        subprocess.run(
            [soffice, "--headless", "--norestore", f"-env:UserInstallation={profile}",
             "--convert-to", csv_filter, "--outdir", tmp, str(workbook)],
            check=True, capture_output=True, timeout=300,
        )
        out = {}
        for path in Path(tmp).glob(f"{workbook.stem}-*.csv"):
            sheet = path.stem[len(workbook.stem) + 1:]
            with path.open(encoding="utf-8", newline="") as f:
                out[sheet] = list(csv.reader(f))
        return out


def color(c, where: str) -> str | None:
    """An ARGB colour as #RRGGBB, or None for the default. Untinted theme colours
    resolve from the workbook's theme; anything else stops the export rather than
    being guessed."""
    if c is None:
        return None
    if c.type == "rgb":
        rgb = str(c.rgb)
        return None if rgb in ("00000000",) else f"#{rgb[-6:].upper()}"
    if c.type == "indexed" and c.indexed in (64, 65):  # system foreground / background
        return None
    if c.type == "theme" and not c.tint and c.theme in THEME:
        return THEME[c.theme]
    detail = f"theme {c.theme}, tint {c.tint}" if c.type == "theme" else f"{c.type} {getattr(c, c.type, '')}"
    fail(f"{where}: colour ({detail}) is not supported; use an RGB or untinted theme colour")


def has_fill(cell) -> bool:
    return bool(cell.fill and cell.fill.fill_type not in (None, "none"))


def is_styled(cell) -> bool:
    return bool(has_fill(cell) or any(
        getattr(cell.border, s).style for s in ("top", "right", "bottom", "left")))


def style_of(cell, where: str) -> dict:
    """The cell's style, keeping only values that differ from the defaults."""
    s: dict = {}
    if cell.number_format and cell.number_format != "General":
        s["nf"] = cell.number_format
    f = cell.font
    font = {k: v for k, v in {
        "name": f.name if f.name and f.name != DEFAULT_FONT_NAME else None,
        "b": bool(f.b) or None, "i": bool(f.i) or None, "u": bool(f.u) or None,
        "sz": float(f.sz) if f.sz and float(f.sz) != 10.0 else None,
        "color": color(f.color, f"{where} font"),
    }.items() if v}
    if font:
        s["font"] = font
    if has_fill(cell):
        if cell.fill.fill_type != "solid":
            fail(f"{where}: {cell.fill.fill_type} fill pattern is not supported")
        s["fill"] = color(cell.fill.fgColor, f"{where} fill")
    a = cell.alignment
    align = {k: v for k, v in {
        "h": a.horizontal if a.horizontal not in (None, "general") else None,
        "v": a.vertical if a.vertical not in (None, "bottom") else None,
        "wrap": True if a.wrap_text else None,
        "indent": int(a.indent) if a.indent else None,
    }.items() if v}
    if align:
        s["align"] = align
    border = {}
    for side, key in (("top", "t"), ("right", "r"), ("bottom", "b"), ("left", "l")):
        b = getattr(cell.border, side)
        if b and b.style:
            border[key] = f"{b.style} {color(b.color, f'{where} border') or '#000000'}"
    if border:
        s["border"] = border
    return s


def negative_red(number_format: str, raw) -> bool:
    """True when the number format paints this (negative) value red, e.g. 0.0%;[Red](0.0%)."""
    if not isinstance(raw, (int, float)) or isinstance(raw, bool) or raw >= 0:
        return False
    sections = number_format.split(";")
    return len(sections) > 1 and "[red]" in sections[1].lower()


def raw_text(v) -> str:
    """How the raw value reads as plain text, to decide whether displayed text differs."""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def widths_and_hidden(ws, columns: list[int]) -> tuple[list[float], set[int]]:
    """Excel widths for the given (1-based) columns, and which positions are hidden."""
    default_w = ws.sheet_format.defaultColWidth or DEFAULT_COL_WIDTH
    by_col: dict[int, float] = {}
    hidden: set[int] = set()
    # Where the file's column definitions overlap, the first one wins, as LibreOffice
    # applies them (definitions are read in file order).
    for dim in ws.column_dimensions.values():
        lo, hi = dim.min or column_index_from_string(dim.index), dim.max or column_index_from_string(dim.index)
        for c in range(lo, hi + 1):
            if dim.width and c not in by_col:
                by_col[c] = round(float(dim.width), 2)
            if dim.hidden:
                hidden.add(c)
    return [by_col.get(c, default_w) for c in columns], {i for i, c in enumerate(columns) if c in hidden}


def row_height(ws, r: int) -> tuple[float, bool]:
    dim = ws.row_dimensions.get(r)
    default_h = ws.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT
    return (round(float(dim.height), 2) if dim is not None and dim.height else default_h), bool(dim is not None and dim.hidden)


class Styles:
    """The deduplicated style table; style 0 is the default and is omitted from cells."""

    def __init__(self) -> None:
        self.table: list[dict] = [{}]
        self.index = {json.dumps({}, sort_keys=True): 0}

    def id(self, style: dict) -> int:
        key = json.dumps(style, sort_keys=True)
        if key not in self.index:
            self.index[key] = len(self.table)
            self.table.append(style)
        return self.index[key]


def cell_entry(sheet, ws_f, ws_v, shown, r, c, pos, styles: Styles, cfg, cf_fills=None) -> dict | None:
    """One exported cell (1-based r, c in the sheet; pos = its position in the tab),
    or None for a blank, unstyled cell. cf_fills: fills painted by colour scales."""
    cell = ws_f.cell(r, c)
    where = f"{sheet}!{cell.coordinate}"
    formula = None
    if cell.data_type == "f":
        formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", None)
        if not isinstance(formula, str):
            fail(f"{where}: unsupported formula type {type(cell.value).__name__}")
        if not formula.startswith("="):
            formula = "=" + formula
    raw = ws_v.cell(r, c).value if formula else cell.value
    if formula and RECALC_FORMULAS:
        again = RECALC_FORMULAS[sheet].cell(r, c).value
        again = again if isinstance(again, str) else getattr(again, "text", again)
        if again != formula and f"={again}" != formula:
            fail(f"{where}: formula differs after recalculation: {formula} vs {again}")
        if raw is None:
            fail(f"{where}: formula {formula} has no value after recalculation")
        if ws_v.cell(r, c).data_type == "e" or (isinstance(raw, str) and raw.startswith("#")):
            fail(f"{where}: formula {formula} recalculates to the error {raw}")
    if formula and raw is None:
        if (sheet, cell.coordinate) not in cfg.get("known_empty_text_formulas", set()):
            fail(f"{where}: formula {formula} has no cached value; open and recalculate the workbook, then save it")
        raw = ""  # known empty-string text formula: shown as an empty cell
    if raw is None and not formula and not is_styled(cell):
        return None

    style = style_of(cell, where)
    if cf_fills and (r, c) in cf_fills:
        style = {**style, "fill": cf_fills[(r, c)]}  # a colour scale paints over the cell's own fill
    if negative_red(cell.number_format, raw):
        style = {**style, "font": {**style.get("font", {}), "color": "#FF0000"}}
    sid = styles.id(style)

    out = {"r": pos[0], "c": pos[1]}
    if sid:
        out["s"] = sid
    if formula:
        out["f"] = formula
    if raw is not None:
        if hasattr(raw, "isoformat"):
            raw = raw.isoformat()
        out["v"] = raw
    shown_row = shown[r - 1] if r - 1 < len(shown) else []
    displayed = shown_row[c - 1] if c - 1 < len(shown_row) else ""
    if raw == "":
        displayed = ""
    if isinstance(raw, str) and ENTITY.search(displayed):
        # Text that literally contains "&amp;" etc. is shown decoded; the raw value keeps it.
        DECODED.append(f"{where}: {', '.join(sorted(set(ENTITY.findall(displayed))))}")
        displayed = html.unescape(displayed)
    if raw is not None and displayed != raw_text(raw):
        out["d"] = displayed
    return out


CFVO_TYPES = {"min", "max", "num", "percent", "percentile"}


def percentile_inc(values: list[float], p: float) -> float:
    """Excel's PERCENTILE.INC, which colour-scale percentile thresholds use."""
    xs = sorted(values)
    k = (len(xs) - 1) * p
    lo = int(k)
    hi = min(lo + 1, len(xs) - 1)
    return xs[lo] + (xs[hi] - xs[lo]) * (k - lo)


def color_scale_fills(sheet, ws_v, rule, sqref) -> dict[tuple[int, int], str]:
    """The fill Excel paints on each numeric cell of a colour-scale rule: thresholds
    from the rule's cfvo list over the whole rule range, then each value linearly
    interpolated between the two surrounding threshold colours, per RGB channel."""
    cs = rule.colorScale
    where = f"{sheet}!{sqref} colour scale"
    if len(cs.cfvo) not in (2, 3) or len(cs.color) != len(cs.cfvo):
        fail(f"{where}: only 2- and 3-colour scales are supported")
    cells = [(r, c) for rng in str(sqref).split()
             for c1, r1, c2, r2 in [range_boundaries(rng)]
             for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    values = {rc: v for rc in cells
              if isinstance(v := ws_v.cell(*rc).value, (int, float)) and not isinstance(v, bool)}
    if not values:
        return {}
    nums = list(values.values())
    lo, hi = min(nums), max(nums)
    stops = []
    for cfvo, col in zip(cs.cfvo, cs.color):
        if cfvo.type not in CFVO_TYPES:
            fail(f"{where}: threshold type {cfvo.type!r} is not supported")
        v = float(cfvo.val) if cfvo.val is not None else None
        t = {"min": lo, "max": hi, "num": v,
             "percent": None if v is None else lo + (hi - lo) * v / 100,
             "percentile": None if v is None else percentile_inc(nums, v / 100)}[cfvo.type]
        rgb = color(col, f"{where} colour")
        if rgb is None or t is None:
            fail(f"{where}: threshold {cfvo.type} has no value or colour")
        stops.append((t, tuple(int(rgb[i:i + 2], 16) for i in (1, 3, 5))))
    out = {}
    for rc, v in values.items():
        if v <= stops[0][0]:
            rgb = stops[0][1]
        elif v >= stops[-1][0]:
            rgb = stops[-1][1]
        else:
            (t0, c0), (t1, c1) = next((a, b) for a, b in zip(stops, stops[1:]) if v <= b[0])
            f = (v - t0) / (t1 - t0) if t1 > t0 else 0.0
            rgb = tuple(round(a + (b - a) * f) for a, b in zip(c0, c1))
        out[rc] = "#" + "".join(f"{x:02X}" for x in rgb)
    return out


def check_sheet(sheet, ws_f, ws_v, bounds=None) -> dict[tuple[int, int], str]:
    """Checks a sheet (or, with bounds, the ranges of a view) for features the preview
    cannot show, and returns the fills that colour scales paint on its cells, keyed by
    1-based (row, column). Conditional formats outside a view's ranges are ignored."""
    def touches(rng: str) -> bool:
        if bounds is None:
            return True
        c1, r1, c2, r2 = range_boundaries(rng)
        return any(c1 <= b[2] and b[0] <= c2 and r1 <= b[3] and b[1] <= r2 for b in bounds)

    rules = [(rule, cf.sqref) for cf in ws_f.conditional_formatting
             if any(touches(rng) for rng in str(cf.sqref).split()) for rule in cf.rules]
    fills: dict[tuple[int, int], str] = {}
    # Lowest priority number wins where rules overlap, so apply it last.
    for rule, sqref in sorted(rules, key=lambda x: x[0].priority or 0, reverse=True):
        if rule.type != "colorScale":
            fail(f"{sheet}!{sqref}: {rule.type} conditional formatting cannot be reproduced by the preview")
        fills.update(color_scale_fills(sheet, ws_v, rule, sqref))
    if getattr(ws_f, "_charts", None) or getattr(ws_f, "_images", None):
        note(f"{sheet} has charts or images; they are not part of the grid")
    return fills


def export_sheet(sheet, ws_f, ws_v, shown, styles: Styles, cfg) -> dict:
    """A whole sheet as one tab."""
    cf_fills = check_sheet(sheet, ws_f, ws_v)
    last_r = last_c = 0
    for row in ws_f.iter_rows():
        for c in row:
            if c.value is not None or is_styled(c):
                last_r, last_c = max(last_r, c.row), max(last_c, c.column)
    rows, cols = min(last_r, MAX_ROWS), min(last_c, MAX_COLS)

    widths, hidden_cols = widths_and_hidden(ws_f, list(range(1, cols + 1)))
    heights, hidden_rows = [], []
    for r in range(1, rows + 1):
        h, hidden = row_height(ws_f, r)
        heights.append(h)
        if hidden:
            hidden_rows.append(r - 1)

    freeze = {"rows": 0, "cols": 0}
    if ws_f.freeze_panes:
        letter, row = coordinate_from_string(ws_f.freeze_panes)
        freeze = {"rows": row - 1, "cols": column_index_from_string(letter) - 1}

    merges = [[m.min_row - 1, m.min_col - 1, min(m.max_row, rows) - 1, min(m.max_col, cols) - 1]
              for m in ws_f.merged_cells.ranges if m.min_row <= rows and m.min_col <= cols]

    cells = []
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            entry = cell_entry(sheet, ws_f, ws_v, shown, r, c, (r - 1, c - 1), styles, cfg, cf_fills)
            if entry:
                cells.append(entry)

    return {
        "name": sheet,
        "rows": rows,
        "cols": cols,
        "truncated": last_r > MAX_ROWS or last_c > MAX_COLS,
        "freeze": freeze,
        "colWidths": widths,
        "rowHeights": heights,
        "hiddenRows": hidden_rows,
        "hiddenCols": sorted(hidden_cols),
        "merges": merges,
        "cells": cells,
    }


def export_view(view, ws_f, ws_v, shown, styles: Styles, cfg) -> dict:
    """A curated view: one or more ranges of a sheet, stacked with a spacer row between
    them. Cells are stored by their position in the view; rowIndex and colIndex give
    the real (zero-based) sheet row and column for each position, null for a spacer."""
    sheet = view["sheet"]
    bounds = [range_boundaries(r) for r in view["ranges"]]  # (min_col, min_row, max_col, max_row), 1-based
    cf_fills = check_sheet(sheet, ws_f, ws_v, bounds)
    first_col, last_col = min(b[0] for b in bounds), max(b[2] for b in bounds)
    columns = list(range(first_col, last_col + 1))

    row_index: list[int | None] = []
    owner: list[tuple | None] = []  # the range each position belongs to
    for i, b in enumerate(bounds):
        if i:
            row_index.append(None)
            owner.append(None)
        for r in range(b[1], b[3] + 1):
            row_index.append(r)
            owner.append(b)

    widths, hidden_cols = widths_and_hidden(ws_f, columns)
    heights, hidden_rows = [], []
    for p, r in enumerate(row_index):
        if r is None:
            heights.append(GAP_ROW_PT)
            continue
        h, hidden = row_height(ws_f, r)
        heights.append(h)
        if hidden:
            hidden_rows.append(p)

    # Merged ranges inside a view range, in view positions. One that a range cuts
    # through cannot be shown as the workbook shows it, so the export stops.
    merges = []
    for m in ws_f.merged_cells.ranges:
        for b in bounds:
            inside = b[0] <= m.min_col and m.max_col <= b[2] and b[1] <= m.min_row and m.max_row <= b[3]
            overlaps = m.min_col <= b[2] and b[0] <= m.max_col and m.min_row <= b[3] and b[1] <= m.max_row
            if inside:
                merges.append([row_index.index(m.min_row), m.min_col - first_col,
                               row_index.index(m.max_row), m.max_col - first_col])
            elif overlaps:
                fail(f"{sheet}!{m.coord}: merged range is partly outside the view range "
                     f"{view['ranges'][bounds.index(b)]}; adjust the range to include or exclude it")

    cells = []
    for p, r in enumerate(row_index):
        if r is None:
            continue
        b = owner[p]
        for c in columns:
            if not b[0] <= c <= b[2]:
                continue  # inside the view's columns but outside this range: left blank
            entry = cell_entry(sheet, ws_f, ws_v, shown, r, c, (p, c - first_col), styles, cfg, cf_fills)
            if entry:
                cells.append(entry)

    return {
        "name": view["name"],
        "sheet": sheet,
        "caption": view["caption"],
        "ranges": view["ranges"],
        "rows": len(row_index),
        "cols": len(columns),
        "rowIndex": [None if r is None else r - 1 for r in row_index],
        "colIndex": [c - 1 for c in columns],
        "truncated": False,
        # A view may freeze its first columns (e.g. the row labels) so they stay put
        # while the values scroll; nothing is frozen unless the view asks for it.
        "freeze": {"rows": 0, "cols": view.get("freeze_cols", 0)},
        "colWidths": widths,
        "rowHeights": heights,
        "hiddenRows": hidden_rows,
        "hiddenCols": sorted(hidden_cols),
        "merges": merges,
        "cells": cells,
    }


def export_project(name: str, cfg: dict) -> None:
    workbook = ROOT / cfg["workbook"]
    output = ROOT / cfg["output"]
    if not workbook.exists():
        fail(f"{name}: workbook not found: {workbook}")
    digest = sha256(workbook)
    source = cfg.get("source")
    if source and digest != source["sha256"]:
        fail(f"{name}: {cfg['workbook']} (sha256 {digest}) does not match its source "
             f"{source['repo']}@{source['commit'][:12]}:{source['path']} (sha256 {source['sha256']})")

    load_theme(workbook)
    wb_f = openpyxl.load_workbook(workbook)
    RECALC_FORMULAS.clear()
    tmp = tempfile.TemporaryDirectory()
    values_from = workbook
    if cfg.get("recalculate"):
        # Values and displayed text from a recalculated copy; formulas and styles from the original.
        values_from = recalculated_copy(workbook, tmp.name)
        RECALC_FORMULAS.update({ws.title: ws for ws in openpyxl.load_workbook(values_from).worksheets})
    wb_v = openpyxl.load_workbook(values_from, data_only=True)
    wanted = cfg.get("sheets") or [v["sheet"] for v in cfg["views"]]
    missing = sorted({s for s in wanted if s not in wb_f.sheetnames})
    if missing:
        fail(f"{name}: requested sheet(s) missing from the workbook: {', '.join(missing)}")
    hidden = sorted({s for s in wanted if wb_f[s].sheet_state != "visible"})
    if hidden:
        fail(f"{name}: requested sheet(s) are hidden in the workbook: {', '.join(hidden)}")

    shown = shown_text(values_from)
    no_csv = sorted({s for s in wanted if s not in shown})
    if no_csv:
        fail(f"{name}: LibreOffice produced no displayed values for: {', '.join(no_csv)}")

    styles = Styles()
    DECODED.clear()
    if "views" in cfg:
        tabs = [export_view(v, wb_f[v["sheet"]], wb_v[v["sheet"]], shown[v["sheet"]], styles, cfg) for v in cfg["views"]]
    else:
        tabs = [export_sheet(s, wb_f[s], wb_v[s], shown[s], styles, cfg) for s in cfg["sheets"]]

    meta = {"file": cfg["workbook"], "sha256": digest}
    if source:
        meta["source"] = source
    if cfg.get("recalculate"):
        meta["values"] = {"from": "LibreOffice recalculation (the workbook stores no cached results)",
                          "libreoffice": libreoffice_version()}
    tmp.cleanup()
    doc = {
        "version": FORMAT_VERSION,
        "workbook": meta,
        "note": ("Read-only snapshot of values recalculated by LibreOffice. Displayed text is exported from "
                 "LibreOffice; nothing recalculates on the page." if cfg.get("recalculate") else
                 "Read-only snapshot. Displayed text is exported from LibreOffice; nothing recalculates."),
        "styles": styles.table,
        "sheets": tabs,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(doc, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    for d in DECODED:
        note(f"decoded HTML entities for display in {d}")
    cells = sum(len(t["cells"]) for t in tabs)
    formulas = sum(1 for t in tabs for c in t["cells"] if "f" in c)
    print(f"wrote {cfg['output']}: {len(tabs)} tabs, {cells} cells, {formulas} formulas, "
          f"{len(styles.table)} styles, {output.stat().st_size} bytes")


def main(argv: list[str]) -> None:
    names = argv or list(PROJECTS)
    unknown = [n for n in names if n not in PROJECTS]
    if unknown:
        fail(f"unknown project(s): {', '.join(unknown)}; choose from {', '.join(PROJECTS)}")
    for n in names:
        export_project(n, PROJECTS[n])


if __name__ == "__main__":
    main(sys.argv[1:])
